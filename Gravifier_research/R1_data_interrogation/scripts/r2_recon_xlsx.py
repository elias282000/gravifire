import openpyxl

for path, sheets_hint in [
    ("/home/elias/Codes/NSAC/psidata/PSI-26/Science Documents/PSI-26_Science Documents_As Run_Test Matrix.xlsx", "asrun"),
    ("/home/elias/Codes/NSAC/psidata/PSI-26/Science Documents/PSI-26_Science Documents_Planned_Test Matrix.xlsx", "planned"),
    ("/home/elias/Codes/NSAC/psidata/PSI-99/Science Documents/PSI-99_Science Documents_SAFFIRE-II_MergeResult-jpeg_v5_Final.xlsx", "merge99"),
    ("/home/elias/Codes/NSAC/psidata/PSI-98/Science Documents/PSI-98_Science Documents_SAFFIRE-1_MergeResult-jpeg_Update_Final_v2.xlsx", "merge98"),
]:
    print(f"\n{'=' * 90}\n{path.split('/')[-1]}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets[:3]:
        print(f"  -- '{ws.title}' dims={ws.calculate_dimension()}")
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
            vals = [str(v)[:26] if v is not None else "" for v in row[:12]]
            print(f"     r{r}: {vals}")
    wb.close()
