import openpyxl

wb = openpyxl.load_workbook(
    "/home/elias/Codes/NSAC/psidata/PSI-26/Science Documents/PSI-26_Science Documents_As Run_Test Matrix.xlsx",
    read_only=True, data_only=True)
ws = wb["as run matrix"]
rows = ws.iter_rows(min_row=3, max_row=3, values_only=True)
hdr = [("" if v is None else str(v).strip()) for v in next(rows)]
for i, h in enumerate(hdr):
    print(f"{i:2d}  {h!r}")
# also sample one data row fully
data = ws.iter_rows(min_row=4, max_row=4, values_only=True)
for i, v in enumerate(next(data)):
    print(f"   r4[{i:2d}] = {v!r}")
wb.close()
